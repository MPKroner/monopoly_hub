# -*- coding: utf-8 -*-
import json
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

with open("C:/Users/Admin/.gemini/antigravity/scratch/monopoly-collection-tracker/editions_seed.json", "r", encoding="utf-8") as f:
    editions = json.load(f)

editions.sort(key=lambda x: x["name"])

# 1. Update CSV
csv_file = "C:/Users/Admin/.gemini/antigravity/scratch/monopoly-collection-tracker/monopoly_collection.csv"
with open(csv_file, "w", encoding="utf-8-sig", newline="") as f:
    writer = csv.writer(f, delimiter=";")
    writer.writerow([
        "ID", "Statut (Possédé / Recherche / Non possédé)", "Nom de l'Édition", "Catégorie", 
        "Année", "Pays / Région", "Éditeur", "Image / Visuel Boîte", "État", "Prix d'achat (€)", 
        "Valeur estimée (€)", "Emplacement / Rangement", "Description & Pions", "Notes personnelles"
    ])
    for ed in editions:
        status_str = "Possédé" if ed.get("status") == "owned" else ("Recherché" if ed.get("status") == "wishlist" else "Non possédé")
        writer.writerow([
            ed["id"],
            status_str,
            ed["name"],
            ed["category"],
            ed["year"],
            ed["country"],
            ed["publisher"],
            ed.get("image_url", ""),
            ed.get("condition", ""),
            ed.get("price", ""),
            ed.get("value", ""),
            ed.get("location", ""),
            ed["description"],
            ed.get("notes", "")
        ])

print(f"CSV updated at {csv_file}")

# 2. Update Excel XLSX
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Collection Monopoly"
ws_stats = wb.create_sheet(title="Statistiques")

header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
owned_fill = PatternFill(start_color="D1E7DD", end_color="D1E7DD", fill_type="solid")
owned_font = Font(name="Segoe UI", size=10, bold=True, color="0F5132")
center_align = Alignment(horizontal="center", vertical="center")
left_align = Alignment(horizontal="left", vertical="center")
thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

headers = [
    "Statut", "Nom de l'Édition", "Catégorie", "Année", "Pays / Région",
    "Éditeur", "Image / Boîte", "État", "Prix d'achat (€)", "Valeur estimée (€)", 
    "Emplacement", "Description & Détails", "Notes personnelles", "ID Système"
]

ws.append(headers)
for cell in ws[1]:
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center_align
    cell.border = thin_border

dv_status = DataValidation(type="list", formula1='"Non possédé,Possédé,Recherché"', allow_blank=False)
ws.add_data_validation(dv_status)

dv_etat = DataValidation(type="list", formula1='"Neuf sous blister,Comme neuf,Très bon état,Bon état,Usé,Incomplet"', allow_blank=True)
ws.add_data_validation(dv_etat)

for row_idx, ed in enumerate(editions, 2):
    status_str = "Possédé" if ed.get("status") == "owned" else ("Recherché" if ed.get("status") == "wishlist" else "Non possédé")
    ws.append([
        status_str,
        ed["name"],
        ed["category"],
        ed["year"],
        ed["country"],
        ed["publisher"],
        ed.get("image_url", ""),
        ed.get("condition", ""),
        ed.get("price", ""),
        ed.get("value", ""),
        ed.get("location", ""),
        ed["description"],
        ed.get("notes", ""),
        ed["id"]
    ])
    
    dv_status.add(f"A{row_idx}")
    dv_etat.add(f"H{row_idx}")
    
    for col_idx in range(1, len(headers) + 1):
        c = ws.cell(row=row_idx, column=col_idx)
        c.font = Font(name="Segoe UI", size=10)
        c.border = thin_border
        if col_idx == 1 and status_str == "Possédé":
            c.fill = owned_fill
            c.font = owned_font
        if col_idx in [1, 4, 8]:
            c.alignment = center_align
        elif col_idx in [9, 10]:
            c.alignment = Alignment(horizontal="right", vertical="center")
            c.number_format = '#,##0.00 €'
        else:
            c.alignment = left_align

column_widths = {
    'A': 16, # Statut
    'B': 42, # Nom
    'C': 26, # Catégorie
    'D': 10, # Année
    'E': 18, # Pays
    'F': 22, # Éditeur
    'G': 25, # Image URL
    'H': 20, # État
    'I': 16, # Prix achat
    'J': 18, # Valeur estimée
    'K': 18, # Emplacement
    'L': 55, # Description
    'M': 30, # Notes
    'N': 22  # ID
}

for col, width in column_widths.items():
    ws.column_dimensions[col].width = width

ws.auto_filter.ref = f"A1:N{len(editions) + 1}"

# Stats Sheet
ws_stats.append(["Indicateur", "Valeur"])
ws_stats[1][0].font = header_font
ws_stats[1][0].fill = header_fill
ws_stats[1][1].font = header_font
ws_stats[1][1].fill = header_fill

stats_data = [
    ("Total des éditions répertoriées", f'=COUNTA(\'Collection Monopoly\'!B2:B{len(editions)+1})'),
    ("Nombre d\'éditions possédées", f'=COUNTIF(\'Collection Monopoly\'!A2:A{len(editions)+1}, "Possédé")'),
    ("Nombre d\'éditions recherchées (Wishlist)", f'=COUNTIF(\'Collection Monopoly\'!A2:A{len(editions)+1}, "Recherché")'),
    ("Pourcentage de possession", f'=B3/B2'),
    ("Total investi en achats (€)", f'=SUM(\'Collection Monopoly\'!I2:I{len(editions)+1})'),
    ("Valeur totale estimée (€)", f'=SUM(\'Collection Monopoly\'!J2:J{len(editions)+1})')
]

for label, formula in stats_data:
    ws_stats.append([label, formula])

ws_stats['B5'].number_format = '0.0%'
ws_stats['B6'].number_format = '#,##0.00 €'
ws_stats['B7'].number_format = '#,##0.00 €'

ws_stats.column_dimensions['A'].width = 38
ws_stats.column_dimensions['B'].width = 20

xlsx_file = "C:/Users/Admin/.gemini/antigravity/scratch/monopoly-collection-tracker/monopoly_collection.xlsx"
wb.save(xlsx_file)
print(f"XLSX updated at {xlsx_file}")
